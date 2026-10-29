import csv, io, json, re

TELPAT = re.compile(r"\(?\d{2}\)?\s*\d{4,5}-?\d{4}")

def separa_extra(extra):
    """Separa um telefone escondido no campo de contato extra."""
    if not extra: return "", ""
    m = TELPAT.search(extra)
    if not m: return "", extra
    tel = m.group(0)
    resto = (extra[:m.start()] + extra[m.end():]).strip(" /·-")
    return tel, resto
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = "/home/user/manutencao-pro/vendas-lavanderia-z7/"

# ---- mensagens, extraidas da pagina (fonte unica da verdade) ----
html = io.open(BASE + "rota-z7.html", encoding="utf-8").read()
bloco = html[html.index("var SCRIPTS = ["):html.index("var SCRIPT_MAP")]
bloco = bloco.replace('" + BASE + "', "Luis Carlos")
pat = re.compile(r'\{k:"([^"]+)", t:"((?:[^"\\]|\\.)*)", via:"((?:[^"\\]|\\.)*)",\s*txt:("(?:[^"\\]|\\.)*")\}')
MSGS = []
for k, t, via, txt in pat.findall(bloco):
    MSGS.append((k, json.loads('"%s"' % t), json.loads('"%s"' % via), json.loads(txt)))
assert len(MSGS) == 11, len(MSGS)

SEG2MSG = {
    "Parceria-chave": "bureau", "Hospedagem": "pousada",
    "Short stay/operador": "gestora", "Airbnb/Temporada": "gestora",
    "Predio de studios": "zelador", "Multiplicador": "gestora",
    "Lar de idosos": "ilpi", "Pet shop": "pet", "Barbearia/Salao": "beleza",
    "Estetica/Spa": "beleza", "Academia": "academia", "Restaurante": "restaurante",
}
OVERRIDE = {"Construtora Convisa": "construtora", "Construtora Viplan": "construtora"}

contatos = list(csv.DictReader(io.open(BASE + "lista-contatos.csv", encoding="utf-8")))

# ---- estilos ----
F = "Arial"
H_FONT = Font(name=F, bold=True, size=10, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor="0E4F4E")
B_FONT = Font(name=F, size=10)
T_FONT = Font(name=F, bold=True, size=14, color="0E4F4E")
S_FONT = Font(name=F, size=9, color="59635E")
PREENCHER = PatternFill("solid", fgColor="FFF2CC")
EXEMPLO = PatternFill("solid", fgColor="EAF2F1")
TOPO = Alignment(vertical="top", wrap_text=True)
LINHA = Border(bottom=Side(style="thin", color="D3D8D1"))

wb = Workbook()

# ================= COMO USAR =================
ws = wb.active; ws.title = "Como usar"
ws.sheet_view.showGridLines = False
linhas = [
    ("Rota Comercial Z7 — planilha de prospecção", "titulo"),
    ("Lavanderia de autosserviço do Posto Zandoná Z7 · R. Ottokar Doerffel, 1489 · Anita Garibaldi, Joinville/SC", "sub"),
    ("", ""),
    ("As abas", "h"),
    ("Contatos — os 62 alvos já levantados, em ordem de distância do posto. A mensagem de cada um é montada sozinha.", ""),
    ("Anfitriões Airbnb — aba em branco para você preencher a partir dos anúncios que achar. Monta a mensagem sozinha também.", ""),
    ("Mensagens — os 11 roteiros, um por tipo de cliente. Editou aqui, muda em todas as linhas das outras abas.", ""),
    ("", ""),
    ("O que você preenche", "h"),
    ("Só as células de fundo amarelo. Todo o resto é calculado.", ""),
    ("Na aba Contatos: Status e Observações.", ""),
    ("Na aba Anfitriões Airbnb: tudo até a coluna Onde achou o telefone.", ""),
    ("", ""),
    ("Sobre o telefone do anfitrião do Airbnb", "h"),
    ("O Airbnb não publica telefone de anfitrião em lugar nenhum do anúncio, nem para quem está logado. O número só", ""),
    ("aparece depois de uma reserva confirmada, e normalmente mascarado. Não adianta procurar na página.", ""),
    ("Os três caminhos que funcionam de verdade:", ""),
    ("1. Pegue o nome do prédio no título do anúncio e procure esse prédio no Google e no Instagram. Quem opera várias", ""),
    ("   unidades quase sempre anuncia fora do Airbnb também — e aí o WhatsApp está publicado.", ""),
    ("2. Vá ao prédio. Portaria, zelador e mural alcançam todos os anfitriões de uma vez.", ""),
    ("3. Procure o título do anúncio no OLX, no Imovelweb e no Instagram: é comum o mesmo imóvel estar nos dois lugares,", ""),
    ("   e no anúncio imobiliário o telefone é público.", ""),
    ("", ""),
    ("Não aborde anfitrião pelo chat do Airbnb oferecendo serviço: é contra as regras da plataforma e queima sua imagem.", "alerta"),
    ("", ""),
    ("A coluna Link WhatsApp só funciona para celular (11 dígitos). Para telefone fixo ela fica vazia de propósito.", "obs"),
    ("Ela abre a conversa; a mensagem fica na coluna ao lado, para você copiar e colar.", "obs"),
]
for i, (txt, tipo) in enumerate(linhas, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    if tipo == "titulo": c.font = T_FONT
    elif tipo == "sub": c.font = S_FONT
    elif tipo == "h": c.font = Font(name=F, bold=True, size=11, color="0E4F4E")
    elif tipo == "alerta": c.font = Font(name=F, bold=True, size=10, color="9A6006")
    elif tipo == "obs": c.font = S_FONT
    else: c.font = B_FONT
ws.column_dimensions["A"].width = 120

# ================= MENSAGENS =================
wm = wb.create_sheet("Mensagens")
cab = ["chave", "tipo de cliente", "canal", "texto da mensagem"]
for j, h in enumerate(cab, 1):
    c = wm.cell(row=1, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = TOPO
for i, (k, t, via, txt) in enumerate(MSGS, start=2):
    for j, v in enumerate([k, t, via, txt], 1):
        c = wm.cell(row=i, column=j, value=v); c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
    wm.row_dimensions[i].height = 118
for col, w in zip("ABCD", [14, 42, 30, 110]):
    wm.column_dimensions[col].width = w
wm.freeze_panes = "A2"

# ================= CONTATOS =================
wc = wb.create_sheet("Contatos", 1)
cab = ["Prioridade", "Segmento", "Nome", "Porte", "Endereço", "Bairro", "Distância do Z7",
       "Telefone", "Telefone 2", "Contato extra", "Roteiro", "Tel 1 (dígitos)", "Tel 2 (dígitos)",
       "Celular p/ WhatsApp", "Link WhatsApp", "Mensagem pronta", "Status", "Observações"]
for j, h in enumerate(cab, 1):
    c = wc.cell(row=1, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = TOPO

ordem = {"ate 1 km": 1, "1-3 km": 2, "3-6 km": 3, "verificar": 4}
contatos.sort(key=lambda r: (ordem.get(r["raio_do_z7"], 9), int(r["prioridade"] or 9), r["nome"]))

for i, r in enumerate(contatos, start=2):
    chave = OVERRIDE.get(r["nome"], SEG2MSG.get(r["segmento"], "pousada"))
    tel2, extra = separa_extra(r["contato_extra"])
    vals = [r["prioridade"], r["segmento"], r["nome"], r["porte"], r["endereco"],
            r["bairro"], r["raio_do_z7"], r["telefone"], tel2, extra, chave]
    for j, v in enumerate(vals, 1):
        c = wc.cell(row=i, column=j, value=v); c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
    LIMPA = 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(%s,"(",""),")","")," ",""),"-","")'
    wc.cell(row=i, column=12, value="=" + LIMPA % ("$H%d" % i))
    wc.cell(row=i, column=13, value="=" + LIMPA % ("$I%d" % i))
    wc.cell(row=i, column=14,
        value='=IF(LEN($L%d)=11,$L%d,IF(LEN($M%d)=11,$M%d,""))' % (i, i, i, i))
    wc.cell(row=i, column=15,
        value='=IF($N%d="","",HYPERLINK("https://wa.me/55"&$N%d,"Abrir conversa"))' % (i, i))
    wc.cell(row=i, column=16,
        value='=SUBSTITUTE(INDEX(Mensagens!$D$2:$D$12,MATCH($K%d,Mensagens!$A$2:$A$12,0)),"[nome]",$C%d)' % (i, i))
    for j in (12, 13, 14, 15, 16):
        c = wc.cell(row=i, column=j); c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
    wc.cell(row=i, column=15).font = Font(name=F, size=10, color="0E4F4E", underline="single")
    for j in (17, 18):
        c = wc.cell(row=i, column=j); c.fill = PREENCHER; c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
    wc.row_dimensions[i].height = 56

dv = DataValidation(type="list",
    formula1='"A contatar,Tentei contato,Em negociação,Fechou,Sem interesse"', allow_blank=True)
wc.add_data_validation(dv); dv.add("Q2:Q%d" % (len(contatos) + 1))
for i in range(2, len(contatos) + 2):
    wc.cell(row=i, column=17, value="A contatar")

for col, w in zip("ABCDEFGHIJKLMNOPQR",
                  [9, 19, 36, 15, 34, 18, 15, 17, 17, 30, 13, 14, 14, 16, 15, 78, 16, 32]):
    wc.column_dimensions[col].width = w
wc.freeze_panes = "D2"
wc.auto_filter.ref = "A1:R%d" % (len(contatos) + 1)

# ================= ANFITRIOES AIRBNB =================
wa = wb.create_sheet("Anfitriões Airbnb", 2)
nota = wa.cell(row=1, column=1,
    value="Preencha as colunas amarelas com o que você achar nos anúncios. O telefone NÃO aparece no Airbnb — veja a aba Como usar. A linha 3 é só um exemplo: apague antes de usar.")
nota.font = Font(name=F, bold=True, size=10, color="9A6006"); nota.alignment = Alignment(wrap_text=True, vertical="center")
wa.merge_cells("A1:L1"); wa.row_dimensions[1].height = 30

cab = ["Anúncio (título)", "Prédio", "Bairro", "Anfitrião", "Nº de anúncios do anfitrião",
       "Link do anúncio", "Telefone (quando achar)", "Onde achou o telefone",
       "Telefone (dígitos)", "Tipo", "Link WhatsApp", "Mensagem pronta", "Status"]
for j, h in enumerate(cab, 1):
    c = wa.cell(row=2, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = TOPO

exemplo = ["Estúdio 08", "a localizar (studios numerados 01 a 08)", "Anita Garibaldi",
           "Sergio", 3, "https://www.airbnb.com.br/rooms/...", "(47) 90000-0000",
           "Instagram do prédio"]
for i in range(3, 33):
    for j, v in enumerate(exemplo if i == 3 else [""] * 8, 1):
        c = wa.cell(row=i, column=j, value=v if i == 3 else None)
        c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
        c.fill = EXEMPLO if i == 3 else PREENCHER
    wa.cell(row=i, column=9,
        value='=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($G%d,"(",""),")","")," ",""),"-","")' % i)
    wa.cell(row=i, column=10,
        value='=IF($I%d="","",IF(LEN($I%d)=11,"celular",IF(LEN($I%d)=10,"fixo","conferir")))' % (i, i, i))
    wa.cell(row=i, column=11,
        value='=IF($J%d="celular",HYPERLINK("https://wa.me/55"&$I%d,"Abrir conversa"),"")' % (i, i))
    wa.cell(row=i, column=12,
        value='=IF($A%d="","",SUBSTITUTE(INDEX(Mensagens!$D$2:$D$12,MATCH("anfitriao",Mensagens!$A$2:$A$12,0)),"[nome]",IF($D%d="",$A%d,$D%d)))' % (i, i, i, i))
    for j in (9, 10, 11, 12):
        c = wa.cell(row=i, column=j); c.font = B_FONT; c.alignment = TOPO; c.border = LINHA
    wa.cell(row=i, column=11).font = Font(name=F, size=10, color="0E4F4E", underline="single")
    c = wa.cell(row=i, column=13); c.fill = PREENCHER; c.font = B_FONT; c.border = LINHA
    wa.row_dimensions[i].height = 56

dv2 = DataValidation(type="list",
    formula1='"A contatar,Tentei contato,Em negociação,Fechou,Sem interesse"', allow_blank=True)
wa.add_data_validation(dv2); dv2.add("M3:M32")
for col, w in zip("ABCDEFGHIJKLM", [34, 32, 18, 18, 14, 30, 20, 24, 16, 10, 15, 78, 16]):
    wa.column_dimensions[col].width = w
wa.freeze_panes = "A3"

out = BASE + "prospeccao-lavanderia-z7.xlsx"
wb.save(out)
print("gravado:", out, "| contatos:", len(contatos), "| mensagens:", len(MSGS))
